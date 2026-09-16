import os

from openpyxl import Workbook


def generar_excel(columns: list[str], rows: list[tuple], sheet_name: str, file_path: str) -> int:
    directory = os.path.dirname(file_path)
    if directory:
        os.makedirs(directory, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31]

    sheet.append(columns)
    for row in rows:
        sheet.append(list(row))

    item_code_index = next(
        (i for i, name in enumerate(columns) if name.lower() == "itemcode"),
        None,
    )

    if item_code_index is not None:
        column_letter = sheet.cell(row=1, column=item_code_index + 1).column_letter
        for row_number in range(2, sheet.max_row + 1):
            sheet[f"{column_letter}{row_number}"].number_format = "@"

    workbook.freeze_panes = "A2"
    workbook.save(file_path)

    return len(rows)
