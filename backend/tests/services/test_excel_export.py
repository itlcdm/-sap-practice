from openpyxl import load_workbook

from app.services.excel_export import generar_excel


def test_generar_excel_writes_header_and_rows(tmp_path):
    file_path = tmp_path / "reporte.xlsx"

    row_count = generar_excel(
        columns=["ItemCode", "ItemName", "Stock"],
        rows=[("A00001", "Producto 1", 10), ("A00002", "Producto 2", 5)],
        sheet_name="ALCONSIT",
        file_path=str(file_path),
    )

    assert row_count == 2

    workbook = load_workbook(file_path)
    sheet = workbook["ALCONSIT"]

    assert [c.value for c in sheet[1]] == ["ItemCode", "ItemName", "Stock"]
    assert [c.value for c in sheet[2]] == ["A00001", "Producto 1", 10]
    assert [c.value for c in sheet[3]] == ["A00002", "Producto 2", 5]


def test_generar_excel_forces_text_format_on_itemcode_column(tmp_path):
    file_path = tmp_path / "reporte.xlsx"

    generar_excel(
        columns=["ItemCode", "Stock"],
        rows=[("00123", 10)],
        sheet_name="ALCONSIT",
        file_path=str(file_path),
    )

    workbook = load_workbook(file_path)
    sheet = workbook["ALCONSIT"]

    assert sheet["A2"].number_format == "@"
